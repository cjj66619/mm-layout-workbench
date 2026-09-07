#!/usr/bin/env python3
"""compose_docx.py — 把项目的 Markdown 章节按官方 Word 模板重新合成 layout/main_layout.docx。

输入（项目根目录 <proj>）：
- 章节：默认 layout/_src/sections（text-normalize 的规范化副本）→ 回退 polish/sections → paper/sections；
- 元数据：paper/paper.yaml（title / keywords / sections 顺序），摘要在 00_abstract.md（`**关键词：**` 行）；
- 模板：仓库固化的 _references/template/hwb_template.docx（封面、摘要页、目录、页眉页脚、样式、编号）。

流程：
1. 解析交叉引用（@fig/@tbl/@eq → 图N/表N/式(N)），PDF 图转 PNG；
2. pandoc（--reference-doc 模板）把正文转成 DOCX，公式为原生 OMML；
3. 后处理：标题套模板自动编号（一、/1.1/1.1.1，参考文献/附录不编号）、列表改成模板 (1)(2) 编号、
   图/表题套模板图注/表注、行间公式三栏居中右编号、代码块装进代码清单表、参考文献套 [1] 样式、
   表格三线表 + 内容列宽 + 12→10.5→9 pt 降级（table_fit.py）；
4. 拼装：模板封面（原样保留，年份由用户自行替换）+ 摘要页（题目/摘要/关键词回填）+ 目录 + 正文/参考文献/附录三节；
5. soffice 渲染 PDF → 回填目录页码 → 再渲染一次。

用法：
    python compose_docx.py <proj> [--src layout/_src/sections] [--out layout/main_layout.docx]
                           [--template hwb_template.docx] [--no-pdf] [--toc-depth 3] [--fig-max-cm 15]
                           [--json layout/compose_report.json]
依赖：pandoc、python-docx；可选 soffice（PDF 与目录页码）、pymupdf（PDF 图转 PNG、目录定位）。
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
SKILLS = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILLS / "table-fit" / "scripts"))
from oxml_utils import (ABSTRACT_NUM_ORDERED, CODE_FONT, M_NS, NUM_ID_REFERENCES, R_NS, SONG, STYLE, TEXT_W,  # noqa: E402
                        TNR, child, clear, cn_number, ensure_style, find_style, find_style_by_name, insert_ordered,
                        p_style, p_text, page_break_para, para, para_props, ppr, run, set_bold, set_fonts,
                        set_numbering, set_pstyle, set_size, tab_run)
from table_fit import fit_all_tables, summarize  # noqa: E402

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE.parents[3] / "_references" / "template" / "hwb_template.docx"
ABSTRACT_NAMES = ("00_abstract.md", "abstract.md", "摘要.md")
ABSTRACT_BEGIN = "HWBABSTRACTBEGIN7F3A"
ABSTRACT_END = "HWBABSTRACTEND7F3A"
EQ_TABLE_MARK = "hwb-equation"
CODE_TABLE_MARK = "hwb-code"
EMU_PER_CM = 360000
LEAD_PHRASE = re.compile(r"^(针对问题[一二三四五六七八九十\d]+[，,：:])")
HEADING_STYLE_LEVEL = {STYLE["h1"]: 1, STYLE["h2"]: 2, STYLE["h3"]: 3,
                       "Heading1": 1, "Heading2": 2, "Heading3": 3}
UNNUMBERED_H1 = ("参考文献", "附录", "附　录", "致谢")

IMAGE_MD = re.compile(r'!\[[^\]]*\]\(([^)\s]+\.pdf)(?:\s+"[^"]*")?\)', re.I)
FIG_DEF = re.compile(r'!\[[^\]]*\]\([^)]*\)\s*\{[^}]*#(fig:[\w:.-]+)[^}]*\}')
TBL_DEF = re.compile(r'^(?:Table|表)?\s*:\s*[^\n]*?\{[^}]*#(tbl:[\w:.-]+)[^}]*\}\s*$', re.M)
EQ_DEF = re.compile(r'\$\$(?:[^$]|\$(?!\$))*\$\$\s*(\{[^}]*#(eq:[\w:.-]+)[^}]*\})?')
REF = re.compile(r'(?<![A-Za-z0-9_@])@((?:fig|tbl|eq):[\w:.-]*[\w])')
ATTR_TAIL = re.compile(r'\s*\{[^}]*#(?:fig|tbl|eq):[\w:.-]+[^}]*\}')
REF_ENTRY = re.compile(r'^\[(\d+)\]\s*')

LUA_FILTER = r"""
function Image(el)
  if el.src:match("%.[pP][dD][fF]$") then
    el.src = el.src:gsub("%.[pP][dD][fF]$", ".png")
  end
  return el
end

local PAGE_BREAK = pandoc.RawBlock("openxml", '<w:p><w:r><w:br w:type="page"/></w:r></w:p>')

function Div(el)
  if el.classes:includes("page-break") or el.classes:includes("pagebreak") then
    return PAGE_BREAK
  end
end

function RawBlock(el)
  if el.format == "latex" and (el.text:match("\\newpage") or el.text:match("\\clearpage")) then
    return PAGE_BREAK
  end
end

local fig_n, tab_n = 0, 0

-- 无题注的图/表不占号，与 resolve_crossrefs 的正文编号保持一致
local function has_caption(cap)
  return cap ~= nil and cap.long ~= nil and #cap.long > 0
end

local function prefix_caption(cap, label)
  local first = cap.long[1]
  if first.t == "Plain" or first.t == "Para" then
    first.content:insert(1, pandoc.Str(label))
    first.content:insert(2, pandoc.Space())
  end
  return cap
end

function Figure(el)
  if not has_caption(el.caption) then return el end
  fig_n = fig_n + 1
  el.caption = prefix_caption(el.caption, "图" .. fig_n)
  return el
end

function Table(el)
  if not has_caption(el.caption) then return el end
  tab_n = tab_n + 1
  el.caption = prefix_caption(el.caption, "表" .. tab_n)
  return el
end

-- 作者手写的标题编号（“一、” / “5.2 ”）去掉，统一由模板样式自动编号
local function strip_manual_number(inlines)
  local first = inlines[1]
  if first == nil or first.t ~= "Str" then
    return inlines
  end
  local s = first.text
  local rest
  if s:match("^[一二三四五六七八九十]+、") then
    rest = s:gsub("^[一二三四五六七八九十]+、", "", 1)
  elseif s:match("^%d%d?%.?$") and inlines[2] ~= nil then
    rest = ""
  elseif s:match("^%d+%.%d+[%.%d]*$") then
    rest = ""
  elseif s:match("^%d+%.%d+[%.%d]*") then
    rest = s:gsub("^%d+%.%d+[%.%d]*%s*", "", 1)
  else
    return inlines
  end
  local out = pandoc.List()
  if rest ~= "" then
    out:insert(pandoc.Str(rest))
  end
  local i = 2
  if rest == "" and inlines[2] ~= nil and inlines[2].t == "Space" then
    i = 3
  end
  while i <= #inlines do
    out:insert(inlines[i])
    i = i + 1
  end
  return out
end

function Header(el)
  el.content = strip_manual_number(el.content)
  return el
end

local eq_labels, eq_n = {}, 0

local function collect_math(el)
  if el.mathtype == "DisplayMath" then
    eq_n = eq_n + 1
    for lab in el.text:gmatch("\\label{([^}]+)}") do
      eq_labels[lab] = eq_n
    end
  end
  return nil
end

local function rewrite_ref(el)
  local ref = el.attributes["reference"]
  if ref and eq_labels[ref] then
    return pandoc.Str("(" .. eq_labels[ref] .. ")")
  end
end

local function drop_empty(el)
  if #el.content == 0 then
    return {}
  end
  for _, x in ipairs(el.content) do
    if not (x.t == "Span" and #x.content == 0) then
      return nil
    end
  end
  return {}
end

return {
  { Math = collect_math },
  { Image = Image, Div = Div, RawBlock = RawBlock, Figure = Figure, Table = Table,
    Header = Header, Link = rewrite_ref, Para = drop_empty, Plain = drop_empty },
}
"""


@dataclass
class FrontMatter:
    title: str = ""
    keywords: list[str] = field(default_factory=list)
    abstract_md: str = ""


@dataclass
class Report:
    warnings: list[str] = field(default_factory=list)
    infos: list[str] = field(default_factory=list)
    counts: dict = field(default_factory=dict)
    tables: list[dict] = field(default_factory=list)
    toc: list[dict] = field(default_factory=list)
    files: dict = field(default_factory=dict)

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)
        print(f"[warn] {msg}")

    def info(self, msg: str) -> None:
        self.infos.append(msg)
        print(f"[info] {msg}")


def sh(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, text=True, capture_output=True, encoding="utf-8", errors="replace", **kw)


# --------------------------------------------------------------------------- 外部程序
def find_exe(name: str, env_var: str, win: list[str] = (), mac: list[str] = ()) -> str | None:
    env = os.environ.get(env_var)
    if env and Path(env).exists():
        return env
    hit = shutil.which(name)
    if hit:
        return hit
    cands = win if platform.system() == "Windows" else (mac if platform.system() == "Darwin" else [])
    for c in cands:
        p = Path(os.path.expandvars(c))
        if p.exists():
            return str(p)
    return None


def find_pandoc() -> str | None:
    return find_exe("pandoc", "PANDOC_BIN",
                    [r"%ProgramFiles%\Pandoc\pandoc.exe", r"%LocalAppData%\Pandoc\pandoc.exe"],
                    ["/opt/homebrew/bin/pandoc", "/usr/local/bin/pandoc"])


def find_soffice() -> str | None:
    return find_exe("soffice", "SOFFICE_BIN",
                    [r"%ProgramFiles%\LibreOffice\program\soffice.exe", r"%ProgramFiles(x86)%\LibreOffice\program\soffice.exe"],
                    ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]) or shutil.which("libreoffice")


# --------------------------------------------------------------------------- 源码收集
def load_meta(paper: Path) -> dict:
    f = paper / "paper.yaml"
    if not f.exists():
        return {}
    text = f.read_text(encoding="utf-8")
    try:
        import yaml  # type: ignore
        data = yaml.safe_load(text) or {}
        return data if isinstance(data, dict) else {}
    except ImportError:
        pass
    data: dict = {}
    key = None
    for line in text.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if line.startswith((" ", "\t", "-")) and key:
            item = line.strip().lstrip("-").strip().strip("\"'")
            data.setdefault(key, [])
            if isinstance(data[key], list):
                data[key].append(item)
            continue
        if ":" in line:
            key, _, val = line.partition(":")
            key, val = key.strip(), val.strip().strip("\"'")
            data[key] = val if val else []
    return data


def pick_sections_dir(proj: Path, src: str | None) -> Path:
    if src:
        p = proj / src
        if not p.is_dir():
            sys.exit(f"[error] 章节目录不存在: {p}")
        return p
    for cand in ("layout/_src/sections", "polish/sections", "paper/sections"):
        p = proj / cand
        if p.is_dir() and any(p.glob("*.md")):
            return p
    sys.exit(f"[error] {proj} 下没有 layout/_src/sections、polish/sections 或 paper/sections")


def original_sections_dir(proj: Path, sec_dir: Path) -> Path:
    """规范化副本的图片相对路径要按原始章节目录解析。"""
    if sec_dir.name == "sections" and sec_dir.parent.name == "_src":
        for cand in ("polish/sections", "paper/sections"):
            p = proj / cand
            if p.is_dir():
                return p
    return sec_dir


def collect_sections(sec_dir: Path, meta: dict) -> tuple[Path | None, list[Path]]:
    order = meta.get("sections")
    files: list[Path] = []
    if isinstance(order, list) and order:
        for name in order:
            p = sec_dir / str(name)
            if not p.suffix:
                p = p.with_suffix(".md")
            if p.exists():
                files.append(p)
            else:
                print(f"[warn] paper.yaml sections 中的 {name} 不存在，跳过")
        extra = sorted(p for p in sec_dir.glob("*.md") if p not in files and not p.name.startswith("_"))
        for p in extra:
            print(f"[warn] {p.name} 未列在 paper.yaml sections 中，按文件名追加到末尾")
        files += extra
    else:
        files = sorted(p for p in sec_dir.glob("*.md") if not p.name.startswith("_"))
    abstract = next((p for p in files if p.name in ABSTRACT_NAMES), None)
    return abstract, [p for p in files if p is not abstract]


def extract_front_matter(meta: dict, abstract_file: Path | None) -> FrontMatter:
    fm = FrontMatter(title=str(meta.get("title") or "").strip())
    kws = meta.get("keywords") or []
    if isinstance(kws, str):
        kws = re.split(r"[;；,，]", kws)
    fm.keywords = [str(k).strip() for k in kws if str(k).strip()]
    if abstract_file is not None:
        text = abstract_file.read_text(encoding="utf-8")
        text = re.sub(r"^---\n.*?\n---\n", "", text, flags=re.S)
        text = re.sub(r"^#+\s*摘\s*要[：:]?\s*$", "", text, flags=re.M)
        m = re.search(r"^\**关键词\**[：:]\**\s*(.+)$", text, flags=re.M)
        if m:
            found = [k.strip() for k in re.split(r"[;；,，]|\u3000+|\s{2,}", m.group(1)) if k.strip()]
            if not fm.keywords:
                fm.keywords = found
            text = text[:m.start()] + text[m.end():]
        fm.abstract_md = text.strip()
    return fm


def resolve_crossrefs(sections: list[Path], rep: Report) -> tuple[dict[str, str], list[str]]:
    numbers: dict[str, str] = {}
    fig_n = tbl_n = eq_n = 0
    raw: list[str] = []
    for sec in sections:
        text = sec.read_text(encoding="utf-8")
        code_free = re.sub(r"```.*?```", lambda m: " " * len(m.group(0)), text, flags=re.S)
        events: list[tuple[int, str, str | None]] = []
        for m in FIG_DEF.finditer(code_free):
            events.append((m.start(), "fig", m.group(1)))
        for m in TBL_DEF.finditer(code_free):
            events.append((m.start(), "tbl", m.group(1)))
        for m in EQ_DEF.finditer(code_free):
            events.append((m.start(), "eq", m.group(2)))
        for _, kind, label in sorted(events):
            if kind == "fig":
                fig_n += 1
                numbers[label] = f"图{fig_n}"
            elif kind == "tbl":
                tbl_n += 1
                numbers[label] = f"表{tbl_n}"
            else:
                eq_n += 1
                if label:
                    numbers[label] = f"式({eq_n})"
        raw.append(text)
    out: list[str] = []
    for text in raw:
        def _sub(m: re.Match) -> str:
            lab = m.group(1)
            if lab in numbers:
                return numbers[lab]
            rep.warn(f"未定义的交叉引用 @{lab}，原样保留")
            return m.group(0)
        text = REF.sub(_sub, text)
        text = ATTR_TAIL.sub("", text)
        out.append(text)
    rep.counts.update(figures_defined=fig_n, tables_defined=tbl_n, equations_defined=eq_n)
    return numbers, out


def convert_pdf_figures(texts: list[str], base: Path, dpi: int, rep: Report) -> None:
    try:
        import pymupdf  # type: ignore
    except ImportError:
        pymupdf = None
    pdftoppm = shutil.which("pdftoppm")
    seen: set[Path] = set()
    for text in texts:
        for m in IMAGE_MD.finditer(text):
            pdf = (base / m.group(1)).resolve()
            if pdf in seen:
                continue
            seen.add(pdf)
            png = pdf.with_suffix(".png")
            if not pdf.exists():
                rep.warn(f"图片不存在: {pdf}")
                continue
            if png.exists() and png.stat().st_mtime >= pdf.stat().st_mtime:
                continue
            if pymupdf is not None:
                doc = pymupdf.open(pdf)
                pix = doc[0].get_pixmap(dpi=dpi)
                pix.set_dpi(dpi, dpi)
                pix.save(png)
                doc.close()
            elif pdftoppm:
                sh([pdftoppm, "-png", "-r", str(dpi), "-singlefile", str(pdf), str(png.with_suffix(""))])
            else:
                rep.warn(f"无 pymupdf/pdftoppm，无法把 {pdf.name} 转成 PNG")


# --------------------------------------------------------------------------- 样式修正（pandoc 自建样式 → 模板口径）
def fix_styles(d) -> None:
    styles_el = d.styles.element
    normal = STYLE["normal"]
    for sid, name in (("BodyText", "Body Text"), ("FirstParagraph", "First Paragraph"), ("Compact", "Compact"),
                      ("BlockText", "Block Text")):
        st = ensure_style(styles_el, sid, name, based_on=normal)
        clear(st, "w:pPr", "w:rPr")
        child(st, "w:basedOn", val=normal)
        if sid == "Compact":
            pp = child(st, "w:pPr")
            para_props(pp, first_line_chars=0, before=0, after=0)
    # 代码：等宽字体、允许换行
    for sid, name in (("SourceCode", "Source Code"),):
        st = ensure_style(styles_el, sid, name, based_on=normal)
        clear(st, "w:pPr", "w:rPr")
        pp = child(st, "w:pPr")
        para_props(pp, first_line_chars=0, before=0, after=0, line=240, line_rule="auto", jc="left")
        child(pp, "w:wordWrap", val="1")
        rp = child(st, "w:rPr")
        set_fonts(rp, CODE_FONT, SONG)
        set_size(rp, 18)
    vc = find_style_by_name(styles_el, "VerbatimChar", "Verbatim Char")
    if vc is not None:
        rp = child(vc, "w:rPr")
        set_fonts(rp, CODE_FONT, SONG)
        clear(rp, "w:sz", "w:szCs", "w:shd", "w:color")
    # 正文默认字体明确写死为 宋体 / Times New Roman
    st = find_style(styles_el, normal)
    if st is not None:
        set_fonts(child(st, "w:rPr"), TNR, SONG)
    # 表题/图题：模板样式 + 关闭自动编号（编号由正文字面给出，与交叉引用一致）
    for sid in (STYLE["table_caption"], STYLE["figure_caption"]):
        st = find_style(styles_el, sid)
        if st is not None:
            pp = child(st, "w:pPr")
            clear(pp, "w:numPr")
            para_props(pp, first_line_chars=0, jc="center")
    fc = find_style(styles_el, STYLE["figure_caption"])
    if fc is not None:
        para_props(child(fc, "w:pPr"), keep_next=False)
    # 目录标题：黑体居中，不带主题色
    toc_h = find_style(styles_el, STYLE["toc_heading"])
    if toc_h is not None:
        rp = child(toc_h, "w:rPr")
        clear(rp, "w:color", "w:rFonts")
        set_fonts(rp, TNR, "黑体")
        pp = child(toc_h, "w:pPr")
        para_props(pp, jc="center", before=120, after=120, line=240, line_rule="auto")


# --------------------------------------------------------------------------- 标题
def heading_level(p) -> int:
    return HEADING_STYLE_LEVEL.get(p_style(p), 0)


def process_headings(body, rep: Report) -> list[tuple[int, str, str]]:
    """标题套模板编号；返回 [(level, 显示文本含编号, 纯标题)] 供目录使用。参考文献/附录及其子标题不编号。"""
    out: list[tuple[int, str, str]] = []
    n1 = n2 = n3 = 0
    unnumbered_block = False
    for p in body.iter(qn("w:p")):
        lvl = heading_level(p)
        if not lvl:
            continue
        text = p_text(p).strip()
        pp = ppr(p)
        clear(pp, "w:numPr", "w:ind", "w:spacing", "w:jc", "w:rPr")
        if lvl == 1:
            unnumbered_block = text.startswith(UNNUMBERED_H1)
        if unnumbered_block:
            set_numbering(p, 0, lvl - 1)
            shown = text
        else:
            if lvl == 1:
                n1 += 1
                n2 = n3 = 0
                shown = f"{cn_number(n1)}、{text}"
            elif lvl == 2:
                n2 += 1
                n3 = 0
                shown = f"{n1}.{n2} {text}"
            else:
                n3 += 1
                shown = f"{n1}.{n2}.{n3} {text}"
        para_props(pp, keep_next=True, keep_lines=True)
        # 标题里不允许出现手写“图1”/“表1”前缀之类的编号残留，保持原文
        out.append((lvl, shown, text))
    rep.counts.update(h1=sum(1 for l, _, _ in out if l == 1), h2=sum(1 for l, _, _ in out if l == 2),
                      h3=sum(1 for l, _, _ in out if l == 3))
    return out


# --------------------------------------------------------------------------- 列表
def process_lists(d, rep: Report) -> int:
    """pandoc 生成的项目符号/编号列表 → 模板 (1) (2) 编号（每个列表从 1 重新开始）。"""
    body = d.element.body
    numbering = d.part.numbering_part.element
    template_num_ids = {int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))
                        if int(n.get(qn("w:numId"))) < 1000}
    next_id = max((int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))), default=0) + 1
    mapping: dict[int, int] = {}
    n = 0
    for p in body.iter(qn("w:p")):
        if heading_level(p) or p_style(p) in (STYLE["table_caption"], STYLE["figure_caption"], STYLE["references"]):
            continue
        np_ = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if np_ is None:
            continue
        nid_el = np_.find(qn("w:numId"))
        ilvl_el = np_.find(qn("w:ilvl"))
        if nid_el is None:
            continue
        nid = int(nid_el.get(qn("w:val")))
        if nid == 0 or nid in template_num_ids:
            continue
        if p.getparent().tag == qn("w:tc"):
            continue
        if nid not in mapping:
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(next_id))
            an = OxmlElement("w:abstractNumId")
            an.set(qn("w:val"), str(ABSTRACT_NUM_ORDERED))
            num.append(an)
            for lvl in range(3):
                ov = OxmlElement("w:lvlOverride")
                ov.set(qn("w:ilvl"), str(lvl))
                so = OxmlElement("w:startOverride")
                so.set(qn("w:val"), "1")
                ov.append(so)
                num.append(ov)
            numbering.append(num)
            mapping[nid] = next_id
            next_id += 1
        ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
        set_numbering(p, mapping[nid], min(ilvl, 2))
        set_pstyle(p, STYLE["ordered_list"])
        pp = ppr(p)
        clear(pp, "w:ind", "w:spacing", "w:contextualSpacing")
        para_props(pp, jc="both", before=0, after=0)
        n += 1
    rep.counts["list_items"] = n
    return n


# --------------------------------------------------------------------------- 图 / 表题
def process_figures(body, fig_max_cm: float, rep: Report) -> int:
    n = 0
    widths: list[float] = []
    max_emu = int(fig_max_cm * EMU_PER_CM)
    for p in list(body.iter(qn("w:p"))):
        st = p_style(p)
        has_drawing = p.find(f".//{qn('w:drawing')}") is not None
        if st in ("Figure", "CaptionedFigure") or (has_drawing and st in ("", "BodyText", "FirstParagraph", STYLE["normal"])):
            if not has_drawing:
                continue
            n += 1
            set_pstyle(p, STYLE["figure"])
            pp = ppr(p)
            clear(pp, "w:ind", "w:spacing")
            para_props(pp, jc="center", keep_next=True, first_line_chars=0, before=120, after=60)
            for ext in p.iter():
                if ext.tag in ("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent",
                               "{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
                    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                    if cx > max_emu:
                        k = max_emu / cx
                        ext.set("cx", str(int(cx * k)))
                        ext.set("cy", str(int(cy * k)))
            ext = p.find(f".//{{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}}extent")
            if ext is not None:
                widths.append(int(ext.get("cx")) / EMU_PER_CM)
        elif st == "ImageCaption":
            set_pstyle(p, STYLE["figure_caption"])
            pp = ppr(p)
            clear(pp, "w:ind", "w:spacing", "w:jc", "w:keepNext")
            set_numbering(p, 0, 0)
            para_props(pp, jc="center", first_line_chars=0)
        elif st == "TableCaption":
            set_pstyle(p, STYLE["table_caption"])
            pp = ppr(p)
            clear(pp, "w:ind", "w:spacing", "w:jc")
            set_numbering(p, 0, 0)
            para_props(pp, jc="center", first_line_chars=0, keep_next=True)
    rep.counts["images"] = n
    if widths:
        rep.counts["figure_width_cm_min"] = round(min(widths), 1)
        rep.counts["figure_width_cm_max"] = round(max(widths), 1)
    return n


# --------------------------------------------------------------------------- 行间公式
def _equation_table(omath_para, number: int):
    tbl = OxmlElement("w:tbl")
    tblpr = child(tbl, "w:tblPr")
    child(tblpr, "w:tblStyle", val=STYLE["equation_table"])
    child(tblpr, "w:tblW", w=TEXT_W, type="dxa")
    child(tblpr, "w:jc", val="center")
    bd = child(tblpr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child(bd, f"w:{side}", val="nil")
    child(tblpr, "w:tblLayout", type="fixed")
    mar = child(tblpr, "w:tblCellMar")
    for side in ("left", "right"):
        child(mar, f"w:{side}", w=0, type="dxa")
    child(tblpr, "w:tblLook", val="0000", firstRow=0, lastRow=0, firstColumn=0, lastColumn=0, noHBand=1, noVBand=1)
    child(tblpr, "w:tblCaption", val=EQ_TABLE_MARK)
    grid = OxmlElement("w:tblGrid")
    tbl.append(grid)
    side_w = 900
    widths = (side_w, TEXT_W - 2 * side_w, side_w)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    tr = OxmlElement("w:tr")
    tbl.append(tr)
    child(child(tr, "w:trPr"), "w:cantSplit")
    for i, w in enumerate(widths):
        tc = OxmlElement("w:tc")
        tcpr = child(tc, "w:tcPr")
        child(tcpr, "w:tcW", w=w, type="dxa")
        child(tcpr, "w:vAlign", val="center")
        p = para()
        pp = ppr(p)
        para_props(pp, first_line_chars=0, before=60, after=60, line=240, line_rule="auto",
                   jc="right" if i == 2 else "center")
        if i == 1:
            p.append(omath_para)
        elif i == 2:
            p.append(run(f"({number})", size=24))
        tc.append(p)
        tr.append(tc)
    return tbl


def process_equations(body, rep: Report) -> int:
    n = 0
    made: list = []
    for omp in list(body.iter(f"{{{M_NS}}}oMathPara")):
        p = omp.getparent()
        if p is None or p.tag != qn("w:p") or p.getparent().tag == qn("w:tc"):
            continue
        n += 1
        tbl = _equation_table(omp, n)
        p.addprevious(tbl)
        made.append(tbl)
        if not p_text(p).strip() and p.find(f".//{qn('w:drawing')}") is None:
            p.getparent().remove(p)
    rep.counts["display_equations"] = n
    return n


def separate_adjacent_tables(body) -> int:
    """相邻两张表之间插一个 1 pt 空段，避免 Word 把它们粘成一张。"""
    n = 0
    for tbl in list(body.iterchildren(qn("w:tbl"))):
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:tbl"):
            spacer = para()
            pp = ppr(spacer)
            para_props(pp, before=0, after=0, line=20, line_rule="exact")
            set_size(child(pp, "w:rPr"), 2)
            tbl.addnext(spacer)
            n += 1
    return n


# --------------------------------------------------------------------------- 代码块
def process_code_blocks(body, rep: Report) -> int:
    """连续的 SourceCode 段落装进单格「代码清单」表（模板附录做法），9 pt 等宽、允许换行、可跨页。"""
    n = 0
    paras = [p for p in body.iter(qn("w:p")) if p_style(p) == "SourceCode" and p.getparent().tag != qn("w:tc")]
    i = 0
    while i < len(paras):
        block = [paras[i]]
        while i + 1 < len(paras) and paras[i + 1] is paras[i].getnext():
            i += 1
            block.append(paras[i])
        i += 1
        n += 1
        tbl = OxmlElement("w:tbl")
        tblpr = child(tbl, "w:tblPr")
        child(tblpr, "w:tblStyle", val=STYLE["code_table"])
        child(tblpr, "w:tblW", w=TEXT_W, type="dxa")
        child(tblpr, "w:jc", val="center")
        child(tblpr, "w:tblLayout", type="fixed")
        mar = child(tblpr, "w:tblCellMar")
        for side in ("left", "right"):
            child(mar, f"w:{side}", w=120, type="dxa")
        child(tblpr, "w:tblLook", val="0000", firstRow=0, lastRow=0, firstColumn=0, lastColumn=0, noHBand=1, noVBand=1)
        child(tblpr, "w:tblCaption", val=CODE_TABLE_MARK)
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
        child(grid, "w:gridCol", w=TEXT_W)
        tr = OxmlElement("w:tr")
        tbl.append(tr)
        tc = OxmlElement("w:tc")
        tcpr = child(tc, "w:tcPr")
        child(tcpr, "w:tcW", w=TEXT_W, type="dxa")
        tr.append(tc)
        block[0].addprevious(tbl)
        for p in block:
            pp = ppr(p)
            clear(pp, "w:ind", "w:pBdr", "w:shd")
            child(pp, "w:wordWrap", val="1")
            para_props(pp, first_line_chars=0, before=0, after=0, line=240, line_rule="auto", jc="left")
            for r in p.iter(qn("w:r")):
                rp = child(r, "w:rPr")
                set_fonts(rp, CODE_FONT, SONG)
                set_size(rp, 18)
            tc.append(p)
        after = para()
        para_props(ppr(after), before=0, after=120, line=240, line_rule="auto")
        tbl.addnext(after)
    rep.counts["code_blocks"] = n
    return n


# --------------------------------------------------------------------------- 参考文献
def process_references(body, rep: Report) -> int:
    n = 0
    in_refs = False
    entries: list = []
    for p in list(body.iterchildren(qn("w:p"))):
        lvl = heading_level(p)
        if lvl == 1:
            in_refs = p_text(p).strip().startswith("参考文献")
            continue
        if in_refs and not lvl and p_text(p).strip():
            entries.append(p)
    literal = [REF_ENTRY.match(p_text(p).strip()) for p in entries]
    sequential = all(m is not None for m in literal) and [int(m.group(1)) for m in literal] == list(range(1, len(entries) + 1))
    for p in entries:
        set_pstyle(p, STYLE["references"])
        pp = ppr(p)
        clear(pp, "w:ind", "w:spacing")
        if sequential:
            _strip_leading(p, REF_ENTRY)
            set_numbering(p, NUM_ID_REFERENCES, 0)
        else:
            set_numbering(p, 0, 0)
        n += 1
    if entries and not sequential:
        rep.warn("参考文献条目编号不是从 [1] 连续递增，保留手写编号（未套自动编号）")
    rep.counts["reference_entries"] = n
    return n


def _strip_leading(p, pattern: re.Pattern) -> None:
    for t in p.iter(qn("w:t")):
        if not (t.text or "").strip():
            continue
        m = pattern.match(t.text.lstrip())
        if m:
            t.text = t.text.lstrip()[m.end():]
        break


# --------------------------------------------------------------------------- 正文段落
def process_body_paragraphs(body) -> None:
    for p in body.iter(qn("w:p")):
        st = p_style(p)
        if st in ("BodyText", "FirstParagraph", "BlockText"):
            if p.getparent().tag == qn("w:tc"):
                continue
            set_pstyle(p, STYLE["normal"])
            pp = ppr(p)
            clear(pp, "w:ind", "w:spacing")
            para_props(pp, first_line_chars=200, jc="both", before=0, after=0)
        elif st == "Compact" and p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None and p.getparent().tag != qn("w:tc"):
            set_pstyle(p, STYLE["normal"])


# --------------------------------------------------------------------------- 模板拼装
def _hash(blob: bytes) -> str:
    return hashlib.sha256(blob).hexdigest()


def _sect_paras(body_el) -> list:
    return [p for p in body_el.iterchildren(qn("w:p")) if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None]


def _remap_rids(el, mapping: dict[str, str]) -> None:
    for node in el.iter():
        for attr in (f"{{{R_NS}}}id", f"{{{R_NS}}}embed", f"{{{R_NS}}}link"):
            v = node.get(attr)
            if v is not None:
                if v not in mapping:
                    raise KeyError(f"模板关系 {v} 无法映射到输出文档")
                node.set(attr, mapping[v])


def build_rid_mapping(tpl_doc, out_doc, needed: set[str]) -> dict[str, str]:
    """模板 document.xml 的 rId → 输出文档 rId：页眉页脚按内容哈希匹配，图片按需新增。"""
    out_by_hash: dict[str, str] = {}
    for rid, rel in out_doc.part.rels.items():
        if rel.is_external:
            continue
        out_by_hash[_hash(rel.target_part.blob)] = rid
    mapping: dict[str, str] = {}
    for rid in needed:
        rel = tpl_doc.part.rels[rid]
        if rel.is_external:
            continue
        h = _hash(rel.target_part.blob)
        if h in out_by_hash:
            mapping[rid] = out_by_hash[h]
        elif rel.reltype.endswith("/image"):
            new_rid, _ = out_doc.part.get_or_add_image(io.BytesIO(rel.target_part.blob))
            mapping[rid] = new_rid
        else:
            new_part = rel.target_part
            mapping[rid] = out_doc.part.relate_to(new_part, rel.reltype)
    return mapping


def _rids_in(el) -> set[str]:
    out: set[str] = set()
    for node in el.iter():
        for attr in (f"{{{R_NS}}}id", f"{{{R_NS}}}embed", f"{{{R_NS}}}link"):
            v = node.get(attr)
            if v is not None:
                out.add(v)
    return out


def _find_para(elems: list, needle: str):
    for el in elems:
        if el.tag == qn("w:p") and needle in p_text(el).replace(" ", "").replace("\u3000", ""):
            return el
    return None


def fill_abstract_page(abs_elems: list, fm: FrontMatter, abstract_paras: list, rep: Report) -> list:
    """模板摘要页：题目行填题目、摘要正文替换、关键词替换；返回新的元素列表。"""
    title_p = _find_para(abs_elems, "题目")
    kw_p = _find_para(list(reversed(abs_elems)), "关键词")
    abs_head = _find_para(abs_elems, "摘要：")
    if abs_head is None:
        abs_head = _find_para(abs_elems, "摘要")
    if title_p is None or kw_p is None or abs_head is None:
        raise SystemExit("[error] 模板摘要页缺少 题目/摘要/关键词 段落，请检查 hwb_template.docx")
    _truncate_after_colon(title_p)
    title_p.append(run(" " + (fm.title or "[论文标题]"), west=TNR, east=SONG, size=32, bold=True))
    _truncate_after_colon(kw_p)
    kw_p.append(run(" " + "  ".join(fm.keywords or ["[关键词]"]), west=TNR, east=SONG, size=24, bold=True))
    # 摘要正文：摘要标题之后、关键词之前的段落全部替换
    i_head, i_kw = abs_elems.index(abs_head), abs_elems.index(kw_p)
    proto = next((el for el in abs_elems[i_head + 1:i_kw] if el.tag == qn("w:p") and p_text(el).strip()), None)
    proto_ppr = copy.deepcopy(proto.find(qn("w:pPr"))) if proto is not None and proto.find(qn("w:pPr")) is not None else None
    new_paras: list = []
    for p in abstract_paras:
        clear(p, "w:pPr")
        if proto_ppr is not None:
            p.insert(0, copy.deepcopy(proto_ppr))
        pp = ppr(p)
        clear(pp, "w:pStyle", "w:numPr", "w:keepNext")
        para_props(pp, first_line_chars=200, jc="both", before=60, after=60)
        text = p_text(p)
        m = LEAD_PHRASE.match(text)
        if m:
            _bold_prefix(p, len(m.group(1)))
        new_paras.append(p)
    if not new_paras:
        rep.warn("没有摘要内容（缺少 00_abstract.md？），摘要页留空")
    spacer = para()
    return abs_elems[:i_head + 1] + new_paras + [spacer] + abs_elems[i_kw:]


def _truncate_after_colon(p) -> None:
    """保留“题 目：”/“关键词：”标签本身，删掉冒号之后的下划线占位。"""
    seen = False
    for r in list(p.findall(qn("w:r"))):
        if seen:
            p.remove(r)
            continue
        for t in r.findall(qn("w:t")):
            txt = t.text or ""
            pos = max(txt.find("："), txt.find(":"))
            if pos >= 0:
                t.text = txt[: pos + 1]
                seen = True
                for later in t.itersiblings(qn("w:t")):
                    r.remove(later)
                break


def _bold_prefix(p, nchars: int) -> None:
    """把段首 nchars 个字符加粗（“针对问题一，”）。"""
    remaining = nchars
    for r in list(p.iter(qn("w:r"))):
        if remaining <= 0:
            break
        ts = r.findall(qn("w:t"))
        if not ts:
            continue
        txt = "".join(t.text or "" for t in ts)
        if len(txt) <= remaining:
            set_bold(child(r, "w:rPr"), True)
            remaining -= len(txt)
        else:
            head = copy.deepcopy(r)
            for t in head.findall(qn("w:t")):
                head.remove(t)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = txt[:remaining]
            head.append(t)
            set_bold(child(head, "w:rPr"), True)
            for t in ts:
                r.remove(t)
            t2 = OxmlElement("w:t")
            t2.set(qn("xml:space"), "preserve")
            t2.text = txt[remaining:]
            r.append(t2)
            r.addprevious(head)
            remaining = 0


def _extract_abstract_paras(body) -> list:
    """取出 pandoc 正文中两个标记段之间的摘要段落（从正文移除）。"""
    paras = list(body.iterchildren())
    begin = next((i for i, p in enumerate(paras) if p.tag == qn("w:p") and ABSTRACT_BEGIN in p_text(p)), None)
    end = next((i for i, p in enumerate(paras) if p.tag == qn("w:p") and ABSTRACT_END in p_text(p)), None)
    if begin is None or end is None:
        return []
    out = paras[begin + 1:end]
    for p in paras[begin:end + 1]:
        body.remove(p)
    return [p for p in out if p.tag == qn("w:p") and (p_text(p).strip() or p.find(f".//{{{M_NS}}}oMath") is not None)]


def build_toc(sdt_tpl, headings: list[tuple[int, str, str]], pages: dict[str, int] | None, depth: int):
    """复用模板目录 SDT 外壳，重建条目：TOC 域 + 静态条目（文本 + 制表符 + 页码）。"""
    sdt = copy.deepcopy(sdt_tpl)
    content = sdt.find(qn("w:sdtContent"))
    paras = list(content.iterchildren(qn("w:p")))
    heading_p = paras[0]
    proto = {}
    for p in paras[1:]:
        st = p_style(p)
        if st in (STYLE["toc1"], STYLE["toc2"], STYLE["toc3"]) and st not in proto:
            proto[st] = copy.deepcopy(p.find(qn("w:pPr")))
    end_p = paras[-1]
    for p in paras[1:-1]:
        content.remove(p)
    entries = [(lvl, shown) for lvl, shown, _ in headings if lvl <= depth]
    for i, (lvl, shown) in enumerate(entries):
        st = STYLE[f"toc{lvl}"]
        p = OxmlElement("w:p")
        pp = copy.deepcopy(proto.get(st)) if proto.get(st) is not None else None
        if pp is None:
            pp = ppr(p)
            child(pp, "w:pStyle", val=st)
        else:
            p.append(pp)
        clear(pp, "w:rPr")
        if i == 0:
            for kind, instr in (("begin", None), (None, ' TOC \\o "1-%d" \\h \\z \\u ' % depth), ("separate", None)):
                r = OxmlElement("w:r")
                if kind:
                    fc = OxmlElement("w:fldChar")
                    fc.set(qn("w:fldCharType"), kind)
                    if kind == "begin":
                        fc.set(qn("w:dirty"), "true")
                    r.append(fc)
                else:
                    it = OxmlElement("w:instrText")
                    it.set(qn("xml:space"), "preserve")
                    it.text = instr
                    r.append(it)
                p.append(r)
        p.append(run(shown, size=21))
        p.append(tab_run())
        pg = (pages or {}).get(shown)
        p.append(run(str(pg) if pg else "", size=21))
        end_p.addprevious(p)
    if not entries:
        for r in list(end_p.findall(qn("w:r"))):
            if r.find(qn("w:fldChar")) is not None:
                end_p.remove(r)
    return sdt


def assemble(out_doc, tpl_doc, fm: FrontMatter, headings, toc_depth: int, rep: Report, pages: dict[str, int] | None = None):
    """把 pandoc 正文接到模板的 封面 / 摘要 / 目录 之后，并按模板分节（正文、参考文献、附录）。"""
    body = out_doc.element.body
    tpl_body = tpl_doc.element.body
    tpl_elems = list(tpl_body.iterchildren())
    sect_paras = _sect_paras(tpl_body)
    final_sect = tpl_body.find(qn("w:sectPr"))
    if len(sect_paras) < 5 or final_sect is None:
        raise SystemExit("[error] 模板分节数与预期（封面/摘要/目录/正文/参考文献/附录）不符")
    idx = [tpl_elems.index(p) for p in sect_paras]
    cover = tpl_elems[: idx[0] + 1]
    abstract = tpl_elems[idx[0] + 1: idx[1] + 1]
    toc_block = tpl_elems[idx[1] + 1: idx[2] + 1]
    body_sect_p, refs_sect_p = sect_paras[3], sect_paras[4]
    sdt_tpl = next((el for el in toc_block if el.tag == qn("w:sdt")), None)
    if sdt_tpl is None:
        raise SystemExit("[error] 模板目录节缺少 TOC 内容控件")

    needed: set[str] = set()
    for el in cover + abstract + toc_block + [body_sect_p, refs_sect_p, final_sect]:
        needed |= _rids_in(el)
    mapping = build_rid_mapping(tpl_doc, out_doc, needed)

    abstract_paras = _extract_abstract_paras(body)
    # pandoc 自己的 sectPr 丢弃
    own_sect = body.find(qn("w:sectPr"))
    if own_sect is not None:
        body.remove(own_sect)
    content = list(body.iterchildren())
    for el in content:
        body.remove(el)

    cover_c = [copy.deepcopy(el) for el in cover]
    abstract_c = [copy.deepcopy(el) for el in abstract]
    abstract_c = fill_abstract_page(abstract_c, fm, abstract_paras, rep)
    toc_c = []
    for el in toc_block:
        if el.tag == qn("w:sdt"):
            toc_c.append(build_toc(el, headings, pages, toc_depth))
        else:
            toc_c.append(copy.deepcopy(el))
    body_sect_c, refs_sect_c, final_c = copy.deepcopy(body_sect_p), copy.deepcopy(refs_sect_p), copy.deepcopy(final_sect)
    # 摘要页从第 1 页起连续编号（封面不显示页码）；摘要节强制另起一页
    for sp in (abstract_c[-1],):
        sect = sp.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")
        if sect is not None:
            child(sect, "w:type", val="nextPage")
            child(sect, "w:pgNumType", start=1)
    for sp in (toc_c[-1], body_sect_c, refs_sect_c):
        sect = sp.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")
        if sect is not None:
            clear(sect, "w:pgNumType")
    clear(final_c, "w:pgNumType")
    for el in cover_c + abstract_c + toc_c + [body_sect_c, refs_sect_c, final_c]:
        _remap_rids(el, mapping)

    # 正文按一级标题切成 正文 / 参考文献 / 附录 三段
    def _h1_start(el, prefixes) -> bool:
        return el.tag == qn("w:p") and heading_level(el) == 1 and p_text(el).strip().startswith(prefixes)
    i_ref = next((i for i, el in enumerate(content) if _h1_start(el, ("参考文献",))), None)
    i_app = next((i for i, el in enumerate(content) if _h1_start(el, ("附录", "附　录"))), None)
    main_part = content[: (i_ref if i_ref is not None else (i_app if i_app is not None else len(content)))]
    refs_part = content[i_ref: (i_app if i_app is not None else len(content))] if i_ref is not None else []
    app_part = content[i_app:] if i_app is not None else []
    if i_ref is None:
        rep.warn("正文没有“参考文献”一级标题，参考文献节省略")
    if i_app is None:
        rep.info("正文没有“附录”一级标题，附录节省略")

    new: list = cover_c + abstract_c + toc_c + main_part + [body_sect_c]
    if refs_part:
        new += refs_part + ([refs_sect_c] if app_part else [])
    new += app_part
    for el in new:
        body.append(el)
    body.append(final_c)
    separate_adjacent_tables(body)
    rep.counts["sections"] = 4 + (1 if refs_part else 0) + (1 if app_part else 0)


# --------------------------------------------------------------------------- PDF / 目录页码
def docx_to_pdf(docx_path: Path, rep: Report) -> Path | None:
    soffice = find_soffice()
    if not soffice:
        rep.info("未找到 soffice，跳过 PDF 渲染与目录页码回填")
        return None
    outdir = docx_path.parent / "_render"
    outdir.mkdir(exist_ok=True)
    r = sh([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)], timeout=600)
    pdf = outdir / docx_path.with_suffix(".pdf").name
    if r.returncode != 0 or not pdf.exists():
        rep.warn(f"soffice 转换失败: {r.stderr[-400:]}")
        return None
    return pdf


def locate_heading_pages(pdf: Path, headings: list[tuple[int, str, str]], rep: Report) -> tuple[dict[str, int], int]:
    """在 PDF 里逐页找标题行，返回 ({显示文本: 显示页码}, 摘要页 PDF 索引)。显示页码：摘要页为 1，之后连续。"""
    try:
        import pymupdf  # type: ignore
    except ImportError:
        rep.info("无 pymupdf，目录页码留空")
        return {}, 0

    def norm(s: str) -> str:
        return re.sub(r"[\s\u3000]+", "", s)

    doc = pymupdf.open(pdf)
    page_lines = [[norm(l) for l in doc[i].get_text().splitlines()] for i in range(doc.page_count)]
    abstract_idx = next((i for i, ls in enumerate(page_lines) if any(l.startswith("摘要：") for l in ls)), 1)
    toc_pages = [i for i, ls in enumerate(page_lines) if i >= abstract_idx and "目录" in ls]
    start = (toc_pages[-1] + 1) if toc_pages else abstract_idx + 1
    found: dict[str, int] = {}
    cur = start
    for _, shown, plain in headings:
        key = norm(plain)
        for pno in range(cur, doc.page_count):
            lines = page_lines[pno]
            hit = any(l.endswith(key) and len(l) - len(key) <= 12 for l in lines)
            if not hit and len(key) >= 18:
                hit = any((a + b).endswith(key) for a, b in zip(lines, lines[1:]))
            if hit:
                found[shown] = pno - abstract_idx + 1
                cur = pno
                break
    doc.close()
    return found, abstract_idx


def fill_toc_pages(docx_path: Path, pages: dict[str, int]) -> None:
    d = docx.Document(str(docx_path))
    sdt = d.element.body.find(qn("w:sdt"))
    if sdt is None:
        return
    for p in sdt.iter(qn("w:p")):
        runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
        if len(runs) < 2:
            continue
        shown = "".join(t.text or "" for t in runs[0].iter(qn("w:t")))
        pg = pages.get(shown)
        if pg:
            t = runs[-1].find(qn("w:t"))
            t.text = str(pg)
    d.save(str(docx_path))


# --------------------------------------------------------------------------- 主流程
def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("proj")
    ap.add_argument("--src", default=None, help="章节目录（相对项目根）")
    ap.add_argument("--out", default="layout/main_layout.docx")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    ap.add_argument("--no-pdf", action="store_true", help="不渲染 PDF（目录页码留空）")
    ap.add_argument("--toc-depth", type=int, default=3, choices=(1, 2, 3))
    ap.add_argument("--fig-max-cm", type=float, default=15.0)
    ap.add_argument("--figure-dpi", type=int, default=300)
    ap.add_argument("--keep-entry", action="store_true", help="保留 pandoc 入口 Markdown 与原始输出（调试）")
    ap.add_argument("--json", default="layout/compose_report.json")
    args = ap.parse_args()

    proj = Path(args.proj).resolve()
    out = (proj / args.out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    template = Path(args.template).resolve()
    if not template.exists():
        sys.exit(f"[error] 模板不存在: {template}")
    pandoc = find_pandoc()
    if not pandoc:
        sys.exit("[error] 未找到 pandoc（或设置 PANDOC_BIN）")

    rep = Report()
    sec_dir = pick_sections_dir(proj, args.src)
    orig_dir = original_sections_dir(proj, sec_dir)
    meta = load_meta(proj / "paper")
    abstract_file, sections = collect_sections(sec_dir, meta)
    fm = extract_front_matter(meta, abstract_file)
    if not fm.title:
        rep.warn("paper/paper.yaml 缺少 title，摘要页题目留占位符")
    numbers, texts = resolve_crossrefs(sections, rep)
    abstract_text = REF.sub(lambda m: numbers.get(m.group(1), m.group(0)), fm.abstract_md)
    convert_pdf_figures(texts + [abstract_text], orig_dir, args.figure_dpi, rep)

    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)
        lua = tdp / "hwb.lua"
        lua.write_text(LUA_FILTER, encoding="utf-8")
        entry = tdp / "_entry.md"
        head = f"{ABSTRACT_BEGIN}\n\n{abstract_text}\n\n{ABSTRACT_END}\n\n" if abstract_text else ""
        entry.write_text(head + "\n\n".join(texts) + "\n", encoding="utf-8")
        raw = tdp / "raw.docx"
        cmd = [pandoc, str(entry), "-f",
               "markdown+tex_math_dollars+raw_tex+implicit_figures+table_captions+pipe_tables+grid_tables+fenced_divs",
               "-t", "docx", "--reference-doc", str(template), "--lua-filter", str(lua),
               "--resource-path", os.pathsep.join(str(p) for p in (orig_dir, proj, proj / "paper", sec_dir)),
               "-o", str(raw)]
        r = sh(cmd, cwd=str(orig_dir))
        if r.returncode != 0:
            if args.keep_entry:
                shutil.copy(entry, out.parent / "_entry.md")
            sys.exit(f"[error] pandoc 失败:\n{r.stderr}")
        for line in r.stderr.splitlines():
            if line.strip():
                rep.warn(f"pandoc: {line.strip()}")
        if args.keep_entry:
            shutil.copy(entry, out.parent / "_entry.md")
            shutil.copy(raw, out.parent / "_pandoc_raw.docx")

        d = docx.Document(str(raw))
        tpl = docx.Document(str(template))
        body = d.element.body
        fix_styles(d)
        headings = process_headings(body, rep)
        process_body_paragraphs(body)
        process_lists(d, rep)
        process_figures(body, args.fig_max_cm, rep)
        process_equations(body, rep)
        process_code_blocks(body, rep)
        process_references(body, rep)
        rep.tables = fit_all_tables(body)
        rep.counts["tables"] = len(rep.tables)
        for w in summarize(rep.tables):
            (rep.warn if w.startswith("WARN") else rep.info)(w[5:])
        assemble(d, tpl, fm, headings, args.toc_depth, rep)
        d.save(str(out))

    pdf = None
    if not args.no_pdf:
        pdf = docx_to_pdf(out, rep)
        if pdf:
            pages, _ = locate_heading_pages(pdf, headings, rep)
            missing = [s for _, s, _ in headings if s not in pages]
            if missing:
                rep.info(f"{len(missing)} 个标题未在 PDF 中定位到页码（目录留空）: {missing[:5]}")
            fill_toc_pages(out, pages)
            rep.toc = [{"level": l, "text": s, "page": pages.get(s)} for l, s, _ in headings]
            pdf = docx_to_pdf(out, rep) or pdf
            final_pdf = out.with_suffix(".pdf")
            shutil.copy(pdf, final_pdf)
            shutil.rmtree(pdf.parent, ignore_errors=True)
            pdf = final_pdf
            try:
                import pymupdf  # type: ignore
                with pymupdf.open(pdf) as doc:
                    rep.counts["pdf_pages"] = doc.page_count
            except ImportError:
                pass
    rep.files = {"docx": str(out), "pdf": str(pdf) if pdf else None, "sections_dir": str(sec_dir),
                 "template": str(template)}
    jpath = proj / args.json
    jpath.parent.mkdir(parents=True, exist_ok=True)
    jpath.write_text(json.dumps({"counts": rep.counts, "warnings": rep.warnings, "infos": rep.infos,
                                 "tables": rep.tables, "toc": rep.toc, "files": rep.files},
                                ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[compose] {out}  pages={rep.counts.get('pdf_pages', '?')}  tables={rep.counts.get('tables')}  "
          f"eq={rep.counts.get('display_equations')}  figs={rep.counts.get('images')}  warn={len(rep.warnings)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
